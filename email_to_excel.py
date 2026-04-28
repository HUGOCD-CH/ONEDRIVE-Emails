#!/usr/bin/env python3
"""
Email to Excel Converter
Reads .txt email files (HTML bodies exported from Outlook) from a folder
and writes an Excel summary with Subject, Date, and full Body content.
Date is taken from the file's last-modified timestamp.

Usage:
    pip install -r requirements.txt
    python email_to_excel.py
"""

import re
import sys
from datetime import datetime
from pathlib import Path
from os.path import getmtime

try:
    from bs4 import BeautifulSoup
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError as exc:
    print(f"Missing dependency: {exc}")
    print("Run:  pip install -r requirements.txt")
    sys.exit(1)


# ── Configuration ──────────────────────────────────────────────────────────────
EMAILS_FOLDER = r"C:\Users\diash2\OneDrive - Medtronic PLC\TEMP\Emails"
OUTPUT_PREFIX = "Email_Summary"


# ── HTML → plain text ──────────────────────────────────────────────────────────
def html_to_text(html: str) -> str:
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["style", "script", "meta", "head", "img"]):
        tag.decompose()
    text = soup.get_text(separator=" ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


# ── Parse one email file ───────────────────────────────────────────────────────
def parse_email_file(path: Path) -> dict:
    content = path.read_text(encoding="utf-8", errors="replace")
    plain = " ".join(html_to_text(content).split())
    file_date = datetime.fromtimestamp(getmtime(path)).strftime("%Y-%m-%d %H:%M")
    return {
        "Subject": path.stem,
        "Date": file_date,
        "Body": plain,
    }


# ── Excel generation ───────────────────────────────────────────────────────────
COLUMNS = [
    ("Subject", 55),
    ("Date",    20),
    ("Body",    80),
]

_HDR_FILL = PatternFill(start_color="00285A", end_color="00285A", fill_type="solid")
_HDR_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
_ROW_FONT = Font(name="Calibri", size=10)
_ALT_FILL = PatternFill(start_color="EEF2F7", end_color="EEF2F7", fill_type="solid")
_WRAP     = Alignment(wrap_text=True, vertical="top")
_CENTER   = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BORDER   = Border(
    bottom=Side(style="thin", color="CCCCCC"),
    right =Side(style="thin", color="CCCCCC"),
)


def generate_excel(emails: list, output_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Email Summary"
    ws.freeze_panes = "A2"

    for col_idx, (header, width) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = _HDR_FONT
        cell.fill      = _HDR_FILL
        cell.alignment = _CENTER
        ws.column_dimensions[cell.column_letter].width = width
    ws.row_dimensions[1].height = 22

    for row_idx, email in enumerate(emails, 2):
        fill = _ALT_FILL if row_idx % 2 == 0 else None
        for col_idx, (header, _) in enumerate(COLUMNS, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=email.get(header, ""))
            cell.font      = _ROW_FONT
            cell.alignment = _WRAP
            cell.border    = _BORDER
            if fill:
                cell.fill = fill

    ws.auto_filter.ref = f"A1:{ws.cell(1, len(COLUMNS)).coordinate}"
    wb.save(output_path)


# ── Main ───────────────────────────────────────────────────────────────────────
def main() -> None:
    folder = Path(EMAILS_FOLDER)

    if not folder.exists():
        print(f"ERROR: Folder not found:\n  {folder}")
        sys.exit(1)

    txt_files = sorted(
        f for f in folder.glob("*.txt")
        if not f.name.startswith(OUTPUT_PREFIX)
    )

    if not txt_files:
        print(f"No .txt email files found in:\n  {folder}")
        sys.exit(0)

    print(f"Found {len(txt_files)} email file(s).\n")

    emails = []
    for path in txt_files:
        try:
            emails.append(parse_email_file(path))
            print(f"  OK  {path.name}")
        except Exception as exc:
            print(f"  ERR {path.name}: {exc}")

    timestamp   = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = folder / f"{OUTPUT_PREFIX}_{timestamp}.xlsx"
    generate_excel(emails, output_path)

    print(f"\nSaved: {output_path}")


if __name__ == "__main__":
    main()
